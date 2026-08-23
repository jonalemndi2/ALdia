"""Contrato de seguridad del parser Eikon (sin base real ni archivo de Jonathan)."""
import io

import pytest
from openpyxl import Workbook

from importadores.eikon import leer_lista, ArchivoEikonInvalido


def _xlsx(cambio=None):
    wb = Workbook(); ws = wb.active; ws.title = "LISTA DE PRECIOS"
    ws.append([]); ws.append([]); ws.append([]); ws.append([])
    ws.append(["Código", "Artículo", "Final Regular USD", "Final Especial USD", "Categoría", "Subcategoría", "IVA", "stock"])
    for i in range(494):
        ws.append([f"EK-{i:04d}", f"Artículo {i}", "10.00", "9.50", "PC", "Partes", 21, 3])
    if cambio: cambio(ws)
    data = io.BytesIO(); wb.save(data); return data.getvalue()


def test_lista_eikon_alfanumerica_y_conversion_half_up():
    filas, errores = leer_lista(_xlsx(), "especial", __import__("decimal").Decimal("1000.005"))
    assert len(filas) == 494 and not errores
    assert filas[0]["sku"] == "EK-0000"
    assert filas[0]["precio_ars_centavos"] == 950005  # Decimal + HALF_UP


@pytest.mark.parametrize("cambio", [
    lambda ws: ws.cell(6, 1).__setattr__("value", "EK-0001"),  # duplicado
    lambda ws: ws.cell(6, 7).__setattr__("value", 7),            # IVA inválido AR
    lambda ws: ws.cell(6, 8).__setattr__("value", -1),           # stock negativo
    lambda ws: ws.cell(6, 3).__setattr__("value", "=1+2"),      # fórmula
])
def test_lista_eikon_rechaza_datos_inseguros(cambio):
    with pytest.raises(ArchivoEikonInvalido):
        leer_lista(_xlsx(cambio), "especial", __import__("decimal").Decimal("1000"))


def test_lista_eikon_exige_494_filas_y_hoja_allowlist():
    with pytest.raises(ArchivoEikonInvalido, match="494"):
        leer_lista(_xlsx(lambda ws: ws.delete_rows(499)), "regular", __import__("decimal").Decimal("1"))


def test_preview_no_escribe_y_confirmacion_solo_deposito(app_cliente):
    """El recorrido HTTP usa la DB temporal de conftest, jamás aldia.db."""
    from database import SessionLocal
    from models import Usuario, StockMercaderia
    from routers.auth import hash_password
    db = SessionLocal()
    try:
        db.add(Usuario(username="deposito-test", password_hash=hash_password("clave-segura-123"),
                       rol="encargado_deposito", debe_cambiar_password=False))
        db.commit()
    finally:
        db.close()
    token = app_cliente.post("/api/auth/login", json={"username": "deposito-test", "password": "clave-segura-123"}).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    r = app_cliente.post("/api/stock/importaciones/preview", headers=headers,
        data={"precio_origen": "especial", "tipo_cambio_ars_por_usd": "1000", "fecha_tc": "2026-08-22", "actualizar_existentes": "false", "cantidad_propia": "0"},
        files={"archivo": ("lista.xlsx", _xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    preview = r.json(); db = SessionLocal()
    try: assert db.query(StockMercaderia).count() == 0
    finally: db.close()
    r = app_cliente.post(f"/api/stock/importaciones/{preview['id']}/confirmar", headers={**headers, "X-Operation-Id": "eikon-test-1"}, json={"confirmar": True, "preview_hash": preview["preview_hash"]})
    assert r.status_code == 200, r.text
    db = SessionLocal()
    try:
        producto = db.query(StockMercaderia).filter_by(sku_proveedor="EK-0000").one()
        assert producto.cantidad == 0 and producto.stock_proveedor == 3
    finally: db.close()
