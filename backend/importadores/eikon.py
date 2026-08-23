"""Lectura cerrada de la lista Eikon.

No conoce la base ni escribe nada.  Mantener este limite hace que subir un
archivo sea una operacion inocua: la unica escritura de catalogo vive en el
endpoint de confirmacion, despues de que la persona revisa el diff.
"""
from __future__ import annotations

from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import io
import json
import zipfile

from openpyxl import load_workbook
from schemas import validar_iva

HOJA = "LISTA DE PRECIOS"
# La lista entregada por Eikon no es estable en la etiqueta de cuatro columnas.
# No hacemos "fuzzy matching": se admiten solamente estas DOS cabeceras completas
# y en el mismo orden. Así una hoja ajena no puede pasar por parecerse un poco.
ENCABEZADOS_CANONICOS = ("Código", "Artículo", "Final Regular USD", "Final Especial USD",
                         "Categoría", "Subcategoría", "IVA", "stock")
ENCABEZADOS_EIKON_V2 = ("Código", "Artículo", "Final Regular", "Final Especial",
                        "Categoría", "Sub Categoría", "% IVA", "stock")
ENCABEZADOS_PERMITIDOS = (ENCABEZADOS_CANONICOS, ENCABEZADOS_EIKON_V2)
MAX_ARCHIVO = 8 * 1024 * 1024
MAX_DESCOMPRIMIDO = 64 * 1024 * 1024
MAX_FILAS = 2_000


class ArchivoEikonInvalido(ValueError):
    pass


def _numero(valor, campo: str) -> Decimal:
    if valor is None or str(valor).strip() == "":
        raise ArchivoEikonInvalido(f"{campo} vacío")
    try:
        n = Decimal(str(valor).strip().replace(",", "."))
    except (InvalidOperation, ValueError):
        raise ArchivoEikonInvalido(f"{campo} no numérico")
    if n < 0:
        raise ArchivoEikonInvalido(f"{campo} no puede ser negativo")
    return n


def _seguridad_zip(datos: bytes) -> None:
    if len(datos) > MAX_ARCHIVO:
        raise ArchivoEikonInvalido("El XLSX supera el máximo de 8 MB")
    try:
        with zipfile.ZipFile(io.BytesIO(datos)) as z:
            total = sum(i.file_size for i in z.infolist())
            if total > MAX_DESCOMPRIMIDO or total > len(datos) * 100:
                raise ArchivoEikonInvalido("XLSX rechazado por expansión ZIP excesiva")
    except zipfile.BadZipFile:
        raise ArchivoEikonInvalido("El archivo no es un XLSX válido")


def leer_lista(datos: bytes, precio_origen: str, tc: Decimal) -> tuple[list[dict], list[dict]]:
    """Devuelve filas válidas y rechazos; no silencia anomalías estructurales."""
    if precio_origen not in {"especial", "regular"}:
        raise ArchivoEikonInvalido("precio_origen debe ser especial o regular")
    if tc <= 0:
        raise ArchivoEikonInvalido("tipo_cambio_ars_por_usd debe ser mayor a cero")
    _seguridad_zip(datos)
    try:
        libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=False)
    except Exception as exc:
        raise ArchivoEikonInvalido("No se pudo abrir el XLSX") from exc
    if HOJA not in libro.sheetnames:
        raise ArchivoEikonInvalido("Solo se admite la hoja LISTA DE PRECIOS")
    hoja = libro[HOJA]
    encabezados = tuple(str(c.value or "").strip() for c in next(hoja.iter_rows(min_row=5, max_row=5)))
    if encabezados[:len(ENCABEZADOS_CANONICOS)] not in ENCABEZADOS_PERMITIDOS:
        raise ArchivoEikonInvalido("Cabecera Eikon inválida o alterada")
    filas, errores, vistos = [], [], set()
    for nro, celdas in enumerate(hoja.iter_rows(min_row=6, values_only=False), 6):
        valores = [c.value for c in celdas[:8]]
        if not any(v not in (None, "") for v in valores):
            continue
        if len(filas) + len(errores) >= MAX_FILAS:
            raise ArchivoEikonInvalido("Demasiadas filas para una importación")
        try:
            if any(getattr(c, "data_type", "") == "f" for c in celdas[:8]):
                raise ArchivoEikonInvalido("Las fórmulas no son admitidas")
            sku = str(valores[0] or "").strip().upper()
            articulo = str(valores[1] or "").strip()
            if not sku or len(sku) > 80:
                raise ArchivoEikonInvalido("Código/SKU inválido")
            if sku in vistos:
                raise ArchivoEikonInvalido("SKU duplicado")
            vistos.add(sku)
            if not articulo or len(articulo) > 200:
                raise ArchivoEikonInvalido("Artículo obligatorio (máximo 200 caracteres)")
            regular, especial = _numero(valores[2], "Final Regular USD"), _numero(valores[3], "Final Especial USD")
            usd = especial if precio_origen == "especial" else regular
            iva = validar_iva(float(_numero(valores[6], "IVA")))
            proveedor_stock = _numero(valores[7], "stock")
            ars_centavos = int((usd * tc * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
            filas.append({"fila": nro, "sku": sku, "producto": articulo,
                          "precio_usd": str(usd), "precio_ars_centavos": ars_centavos,
                          "categoria": str(valores[4] or "").strip()[:120],
                          "subcategoria": str(valores[5] or "").strip()[:120],
                          "iva": iva, "stock_proveedor": float(proveedor_stock)})
        except (ArchivoEikonInvalido, ValueError) as exc:
            errores.append({"fila": nro, "error": str(exc)})
    # La lista v2 conocida tiene 494 productos.  La validación evita importar
    # una hoja desplazada o parcialmente filtrada como si fuese el catálogo.
    if len(filas) != 494:
        raise ArchivoEikonInvalido(f"Se esperaban 494 artículos Eikon y se encontraron {len(filas)}")
    return filas, errores


def hash_preview(filas: list[dict], politica: dict) -> str:
    cuerpo = json.dumps({"filas": filas, "politica": politica}, sort_keys=True,
                        separators=(",", ":"), ensure_ascii=False).encode()
    return hashlib.sha256(cuerpo).hexdigest()
